# -*- coding: utf-8 -*-
"""Generate the WLAN inspection report (Lite Edition) in MD and DOCX format.

Usage:
    python gen_report.py --data-file reports/Inspection_Data_<site>_<time>.json \
        [--data-file ...] --output-dir reports/ [--lang en|zh]

The input format is defined in references/data-format.md. This script only
formats: it computes ratings and statistics from the numbers in the data file
and lays out the six report chapters. All narrative text comes from the
"analysis" block written by the LLM.

Language: --lang, else meta.language, else "en". Supported: en, zh.
With more than one --data-file, one report per site is written plus a
multi-site summary (Inspection_Summary_<timestamp>.md).

The DOCX file needs python-docx. If it is not installed, the MD reports are
still written and the script exits with code 3.
"""
import argparse
import datetime
import json
import os
import re
import sys

STRINGS = {
    "en": {
        "title": "Wireless Network O&M Inspection Report (Lite Edition)",
        "summary_title": "Wireless Network O&M Inspection Summary (Lite Edition)",
        "ok": "🟢 Normal", "warn": "⚠️ Needs attention", "crit": "🔴 Critical", "nodata": "⚪ No data",
        "na": "n/a", "none": "None",
        "cover_site": "Site", "cover_region": "Region", "cover_time": "Inspection time",
        "cover_tool": "Inspection tool", "default_tool": "Cloudnet MCP",
        "h_exec": "1. Executive Summary", "h_health": "2. Health Assessment",
        "h_ac": "2.1 AC Health", "h_dist": "2.2 Problem Distribution",
        "h_access": "2.3 Access Success Rate", "h_ap": "2.4 AP Online Rate",
        "h_detail": "3. Detailed Problem Analysis", "h_high": "3.1 High-risk Problems",
        "h_medium": "3.2 Medium-risk Problems", "h_low": "3.3 Low-risk Problems",
        "h_prio": "4. Problem Handling Priority", "h_rem": "5. Remediation Plan",
        "h_short": "5.1 Short Term", "h_mid": "5.2 Medium Term", "h_long": "5.3 Long Term",
        "h_concl": "6. Inspection Conclusion", "h_risks": "Key Risks",
        "h_trends": "Trend Recommendations", "h_notes": "Data Notes",
        "col_metric": "Metric", "col_value": "Value", "col_rating": "Rating", "col_item": "Item",
        "col_check": "Check item", "col_threshold": "Threshold (Normal / Attention)",
        "col_cat": "Problem category", "col_occ": "Occurrences", "col_sub": "Sub-types",
        "col_field": "Field", "col_prio": "Priority", "col_basis": "Basis", "col_action": "Action",
        "col_type": "Problem type", "col_category": "Category", "col_status": "Status",
        "col_desc": "Description", "col_sugg": "Suggestion",
        "row_ac": "AC {name} CPU / Memory / Disk", "row_ac_health": "AC health",
        "no_ac": "No AC inspected", "row_access": "Access success rate (lowest)",
        "samples": "{n} samples", "no_samples": "No samples", "not_collected": "Not collected",
        "row_ap": "AP online rate", "ap_val": "{rate} ({on} online / {total} total)",
        "row_problems": "Problems recorded", "row_reasoned": "Reasoned problems",
        "reasoned_val": "{h} high-risk / {m} medium-risk", "row_overall": "Overall",
        "no_ac_p": "No AC was inspected for this site.",
        "dev_name": "Device name", "model": "Model", "sn": "Serial number", "dev_type": "Device type",
        "role": "Role", "status": "Status", "online": "Online", "offline": "Offline",
        "soft_ver": "Software version", "address": "Device address",
        "cpu": "CPU", "memory": "Memory", "disk": "Disk",
        "throughput": "Live throughput: {up} Kbps up / {down} Kbps down.",
        "window": "Window: {start} to {end} ({tz}).",
        "no_problems": "No problems were recorded in any category during the window.",
        "dist_conflict": "The API returned non-identical duplicate data for: {types}. The highest count per category is shown.",
        "lowest": "Lowest success rate", "lowest_time": "Time of lowest rate",
        "below98": "Periods below 98%", "n_samples": "Samples in window",
        "total_aps": "Total APs", "ap_online": "Online", "ap_offline": "Offline",
        "online_rate": "Online rate", "offline_list": "Offline AP list",
        "maybe_partial": " (list may be incomplete)",
        "no_level": "None. Cloudnet problem reasoning returned no items with alarm level {lvl}.",
        "reasoning_missing": "Cloudnet problem reasoning was not collected.",
        "reasoning_basis": "Cloudnet reasoning, alarm level {lvl}",
        "no_handling": "No items require handling.",
        "no_prio": "No priorities were provided; review the ratings in chapter 2.",
        "no_actions_ok": "No actions required.", "no_actions": "No actions listed.",
        "overall_p": "Overall assessment: {rating}. {text}",
        "none_identified": "None identified.",
        "sum_col_site": "Site", "sum_col_ac": "AC health", "sum_col_access": "Access (lowest)",
        "sum_col_ap": "AP online", "sum_col_risk": "High / Medium", "sum_col_overall": "Overall",
        "sum_col_report": "Report", "sum_generated": "Generated",
    },
    "zh": {
        "title": "无线网络运维巡检报告（精简版）",
        "summary_title": "无线网络运维巡检汇总（精简版）",
        "ok": "🟢 正常", "warn": "⚠️ 需关注", "crit": "🔴 严重", "nodata": "⚪ 无数据",
        "na": "无", "none": "无",
        "cover_site": "场所", "cover_region": "分支", "cover_time": "巡检时间",
        "cover_tool": "巡检工具", "default_tool": "Cloudnet MCP",
        "h_exec": "1. 执行摘要", "h_health": "2. 健康度评估",
        "h_ac": "2.1 AC 健康度", "h_dist": "2.2 问题分布",
        "h_access": "2.3 接入成功率", "h_ap": "2.4 AP 在线率",
        "h_detail": "3. 问题详细分析", "h_high": "3.1 高风险问题",
        "h_medium": "3.2 中风险问题", "h_low": "3.3 低风险问题",
        "h_prio": "4. 问题处理优先级", "h_rem": "5. 整改计划",
        "h_short": "5.1 短期", "h_mid": "5.2 中期", "h_long": "5.3 长期",
        "h_concl": "6. 巡检结论", "h_risks": "主要风险",
        "h_trends": "趋势建议", "h_notes": "数据说明",
        "col_metric": "指标", "col_value": "数值", "col_rating": "评级", "col_item": "项目",
        "col_check": "检查项", "col_threshold": "阈值（正常 / 关注）",
        "col_cat": "问题大类", "col_occ": "次数", "col_sub": "子类",
        "col_field": "字段", "col_prio": "优先级", "col_basis": "依据", "col_action": "处理措施",
        "col_type": "问题类型", "col_category": "分类", "col_status": "状态",
        "col_desc": "描述", "col_sugg": "建议",
        "row_ac": "AC {name} CPU / 内存 / 磁盘", "row_ac_health": "AC 健康度",
        "no_ac": "未巡检 AC", "row_access": "接入成功率（最低）",
        "samples": "{n} 个采样", "no_samples": "无采样", "not_collected": "未采集",
        "row_ap": "AP 在线率", "ap_val": "{rate}（在线 {on} / 总数 {total}）",
        "row_problems": "问题总数", "row_reasoned": "推理问题",
        "reasoned_val": "高风险 {h} / 中风险 {m}", "row_overall": "总体",
        "no_ac_p": "该场所未巡检 AC。",
        "dev_name": "设备名称", "model": "型号", "sn": "序列号", "dev_type": "设备类型",
        "role": "角色", "status": "状态", "online": "在线", "offline": "离线",
        "soft_ver": "软件版本", "address": "设备地址",
        "cpu": "CPU", "memory": "内存", "disk": "磁盘",
        "throughput": "实时速率：上行 {up} Kbps / 下行 {down} Kbps。",
        "window": "统计时段：{start} 至 {end}（{tz}）。",
        "no_problems": "统计时段内各类问题均未发生。",
        "dist_conflict": "接口对以下类别返回了不一致的重复数据：{types}。表中取每类的最大值。",
        "lowest": "最低成功率", "lowest_time": "最低值出现时间",
        "below98": "低于 98% 的时段数", "n_samples": "采样点数",
        "total_aps": "AP 总数", "ap_online": "在线", "ap_offline": "离线",
        "online_rate": "在线率", "offline_list": "离线 AP 列表",
        "maybe_partial": "（列表可能不完整）",
        "no_level": "无。Cloudnet 问题推理未返回告警级别为 {lvl} 的问题。",
        "reasoning_missing": "未采集到 Cloudnet 问题推理数据。",
        "reasoning_basis": "Cloudnet 问题推理，告警级别 {lvl}",
        "no_handling": "无需处理的问题。",
        "no_prio": "未提供优先级，请参考第 2 章评级。",
        "no_actions_ok": "无需整改。", "no_actions": "未列出措施。",
        "overall_p": "总体评估：{rating}。{text}",
        "none_identified": "未发现。",
        "sum_col_site": "场所", "sum_col_ac": "AC 健康度", "sum_col_access": "接入（最低）",
        "sum_col_ap": "AP 在线", "sum_col_risk": "高 / 中风险", "sum_col_overall": "总体",
        "sum_col_report": "报告", "sum_generated": "生成时间",
    },
}

OK, WARN, CRIT, NODATA = "ok", "warn", "crit", "nodata"
SEVERITY = {OK: 0, NODATA: 0, WARN: 1, CRIT: 2}


# ---------------------------------------------------------------- ratings
def rate_usage(value, normal, attention):
    """Lower is better (CPU / memory / disk)."""
    if value is None:
        return NODATA
    return OK if value <= normal else WARN if value <= attention else CRIT


def rate_rate(value):
    """Higher is better (access success rate / AP online rate)."""
    if value is None:
        return NODATA
    return OK if value >= 98 else WARN if value >= 95 else CRIT


def worst(ratings):
    ratings = [r for r in ratings if r is not None]
    if not ratings:
        return NODATA
    # A real rating beats "no data" when severities tie at 0.
    return max(ratings, key=lambda r: (SEVERITY[r], r != NODATA))


def pct(value, na):
    if value is None:
        return na
    return ("%d%%" % value) if float(value).is_integer() else ("%.1f%%" % value)


# ---------------------------------------------------------------- data prep
def load(path):
    with open(path, encoding="utf-8-sig") as f:
        data = json.load(f)
    missing = [k for k in ("meta", "analysis") if k not in data]
    if missing:
        sys.exit("[gen_report] ERROR: %s is missing required keys: %s "
                 "(see references/data-format.md)" % (path, ", ".join(missing)))
    for k in ("site_name", "inspection_time"):
        if not data["meta"].get(k):
            sys.exit("[gen_report] ERROR: %s: meta.%s is required" % (path, k))
    return data


def dedupe_problems(items):
    """The API may return each category more than once, sometimes with different
    counts. Keep the highest count per type and report the types that differed."""
    merged, seen, conflicts = {}, {}, []
    for p in items or []:
        key = p.get("type") or p.get("name")
        times = p.get("times") or 0
        seen.setdefault(key, set()).add(times)
        if key not in merged or times > (merged[key].get("times") or 0):
            merged[key] = p
    for key, values in seen.items():
        if len(values) > 1:
            conflicts.append(merged[key].get("name") or key)
    return list(merged.values()), conflicts


def compute(data):
    c = {}
    acs = data.get("ac") or []
    for ac in acs:
        ac["model"] = (ac.get("model") or "").replace("\u200b", "").strip()
        ac["_ratings"] = {
            "cpu": rate_usage(ac.get("cpu"), 50, 70),
            "memory": rate_usage(ac.get("memory"), 70, 85),
            "disk": rate_usage(ac.get("disk"), 70, 85),
        }
    c["ac_rating"] = worst([r for ac in acs for r in ac["_ratings"].values()]) if acs else NODATA

    dist = data.get("problem_distribution")
    c["dist_collected"] = dist is not None
    problems, c["dist_conflicts"] = dedupe_problems(dist)
    c["problems"] = problems
    c["problem_total"] = sum(p.get("times") or 0 for p in problems)
    c["top5"] = sorted([p for p in problems if (p.get("times") or 0) > 0],
                       key=lambda p: -(p.get("times") or 0))[:5]

    acc = data.get("access_success")
    c["access_collected"] = acc is not None
    samples = (acc or {}).get("samples") or []
    if samples:
        low = min(samples, key=lambda s: s["rate"])
        c["access_min"], c["access_min_time"] = low["rate"], low.get("time", "")
        c["access_below98"] = sum(1 for s in samples if s["rate"] < 98)
    else:
        c["access_min"], c["access_min_time"], c["access_below98"] = None, "", 0
    c["access_samples"] = len(samples)
    c["access_rating"] = rate_rate(c["access_min"])

    ap = data.get("ap")
    c["ap_collected"] = ap is not None
    total = (ap or {}).get("total") or 0
    c["ap_rate"] = round(((ap or {}).get("online") or 0) * 100.0 / total, 1) if total else None
    c["ap_rating"] = rate_rate(c["ap_rate"])

    reasoning = data.get("reasoning")
    c["reasoning_collected"] = reasoning is not None
    items = (reasoning or {}).get("items") or []
    c["high"] = [i for i in items if (i.get("alarm_level") or 0) >= 3]
    c["medium"] = [i for i in items if (i.get("alarm_level") or 0) == 2]
    c["low"] = [i for i in items if (i.get("alarm_level") or 0) < 2]

    if not c["reasoning_collected"]:
        risk = NODATA
    else:
        risk = CRIT if c["high"] else WARN if c["medium"] else OK
    c["risk_rating"] = risk
    c["problem_rating"] = (OK if c["problem_total"] == 0 else WARN) if c["dist_collected"] else NODATA
    c["overall"] = worst([c["ac_rating"], c["access_rating"], c["ap_rating"], risk])
    return c


# ---------------------------------------------------------------- content
def build_blocks(data, c, S):
    """Return the report as a list of (kind, payload) blocks shared by MD and DOCX.

    kinds: h1, h2, p, bullets (list of str), table ((headers, rows))
    """
    meta, an = data["meta"], data["analysis"]
    ap = data.get("ap") or {}
    acs = data.get("ac") or []
    R = lambda key: S[key]  # rating key -> label
    P = lambda v: pct(v, S["na"])
    b = []
    add = lambda kind, payload: b.append((kind, payload))

    # 1. Executive summary
    add("h1", S["h_exec"])
    if an.get("executive_summary"):
        add("p", an["executive_summary"])
    rows = []
    for ac in acs:
        rows.append([S["row_ac"].format(name=ac.get("name", "")),
                     "%s / %s / %s" % (P(ac.get("cpu")), P(ac.get("memory")), P(ac.get("disk"))),
                     R(worst(list(ac["_ratings"].values())))])
    if not acs:
        rows.append([S["row_ac_health"], S["no_ac"], R(NODATA)])
    if not c["access_collected"]:
        access_val = S["not_collected"]
    elif c["access_samples"]:
        access_val = "%s (%s)" % (P(c["access_min"]), S["samples"].format(n=c["access_samples"]))
    else:
        access_val = S["no_samples"]
    rows += [
        [S["row_access"], access_val, R(c["access_rating"])],
        [S["row_ap"], S["ap_val"].format(rate=P(c["ap_rate"]), on=ap.get("online", 0),
                                         total=ap.get("total", 0))
         if c["ap_collected"] else S["not_collected"], R(c["ap_rating"])],
        [S["row_problems"], str(c["problem_total"]) if c["dist_collected"] else S["not_collected"],
         R(c["problem_rating"])],
        [S["row_reasoned"], S["reasoned_val"].format(h=len(c["high"]), m=len(c["medium"]))
         if c["reasoning_collected"] else S["not_collected"], R(c["risk_rating"])],
        [S["row_overall"], "", R(c["overall"])],
    ]
    add("table", ([S["col_metric"], S["col_value"], S["col_rating"]], rows))

    # 2. Health assessment
    add("h1", S["h_health"])
    add("h2", S["h_ac"])
    if not acs:
        add("p", S["no_ac_p"])
    for ac in acs:
        info = [
            [S["dev_name"], ac.get("name", "")],
            [S["model"], ac.get("model", "")],
            [S["sn"], ac.get("sn", "")],
            [S["dev_type"], ac.get("type", "")],
        ]
        if ac.get("role"):
            info.append([S["role"], ac["role"]])
        info += [
            [S["status"], S["online"] if ac.get("online") else S["offline"]],
            [S["soft_ver"], ac.get("soft_ver", "")],
            [S["address"], ac.get("address", "") or S["na"]],
        ]
        add("table", ([S["col_item"], S["col_value"]], info))
        r = ac["_ratings"]
        add("table", ([S["col_check"], S["col_value"], S["col_threshold"], S["col_rating"]], [
            [S["cpu"], P(ac.get("cpu")), "≤50% / ≤70%", R(r["cpu"])],
            [S["memory"], P(ac.get("memory")), "≤70% / ≤85%", R(r["memory"])],
            [S["disk"], P(ac.get("disk")), "≤70% / ≤85%", R(r["disk"])],
        ]))
        if ac.get("speed_up_kbps") is not None or ac.get("speed_down_kbps") is not None:
            add("p", S["throughput"].format(up=ac.get("speed_up_kbps", S["na"]),
                                            down=ac.get("speed_down_kbps", S["na"])))
        if ac.get("note"):
            add("p", ac["note"])

    add("h2", S["h_dist"])
    tw = meta.get("time_window") or {}
    if tw:
        add("p", S["window"].format(start=tw.get("start", ""), end=tw.get("end", ""),
                                    tz=tw.get("timezone", "")))
    if not c["dist_collected"]:
        add("p", S["not_collected"])
    elif c["top5"]:
        rows = []
        for i, p in enumerate(c["top5"], 1):
            subs = ", ".join("%s (%d)" % (s.get("name", ""), s.get("times", 0))
                             for s in sorted(p.get("sub_types") or [], key=lambda s: -(s.get("times") or 0))
                             if (s.get("times") or 0) > 0)
            rows.append([str(i), p.get("name", p.get("type", "")), str(p.get("times", 0)), subs or "-"])
        add("table", (["#", S["col_cat"], S["col_occ"], S["col_sub"]], rows))
    else:
        add("p", S["no_problems"])
    if c["dist_conflicts"]:
        add("p", S["dist_conflict"].format(types=", ".join(c["dist_conflicts"])))

    add("h2", S["h_access"])
    add("table", ([S["col_item"], S["col_value"]], [
        [S["lowest"], P(c["access_min"]) if c["access_samples"] else access_val],
        [S["lowest_time"], c["access_min_time"] or S["na"]],
        [S["below98"], str(c["access_below98"])],
        [S["n_samples"], str(c["access_samples"])],
        [S["col_rating"], R(c["access_rating"])],
    ]))

    add("h2", S["h_ap"])
    if not c["ap_collected"]:
        add("p", S["not_collected"])
    else:
        offline = ap.get("offline_aps") or []
        offline_text = ", ".join(offline) if offline else S["none"]
        if ap.get("list_complete") is False:
            offline_text += S["maybe_partial"]
        add("table", ([S["col_item"], S["col_value"]], [
            [S["total_aps"], str(ap.get("total", 0))],
            [S["ap_online"], str(ap.get("online", 0))],
            [S["ap_offline"], str(ap.get("offline", 0))],
            [S["online_rate"], P(c["ap_rate"])],
            [S["offline_list"], offline_text],
            [S["col_rating"], R(c["ap_rating"])],
        ]))

    # 3. Detailed problem analysis
    add("h1", S["h_detail"])
    for title, items, level in ((S["h_high"], c["high"], 3), (S["h_medium"], c["medium"], 2)):
        add("h2", title)
        if not c["reasoning_collected"]:
            add("p", S["reasoning_missing"])
            continue
        if not items:
            add("p", S["no_level"].format(lvl=level))
            continue
        for it in items:
            add("table", ([S["col_field"], S["col_value"]], [
                [S["col_type"], it.get("type", "")],
                [S["col_category"], it.get("category", "")],
                [S["col_occ"], str(it.get("count", ""))],
                [S["col_status"], it.get("status", "")],
                [S["col_desc"], it.get("description", "")],
                [S["col_sugg"], it.get("suggestion", "")],
            ]))
    if c["low"]:
        add("h2", S["h_low"])
        add("table", ([S["col_type"], S["col_category"], S["col_occ"], S["col_status"]],
                      [[i.get("type", ""), i.get("category", ""), str(i.get("count", "")), i.get("status", "")]
                       for i in c["low"]]))

    # 4. Priority
    add("h1", S["h_prio"])
    pr = an.get("priorities") or []
    if not pr:
        # Fallback so a risky site never reads "nothing to do": derive P0/P1
        # straight from the reasoning items.
        pr = [{"priority": "P0" if it in c["high"] else "P1", "item": it.get("type", ""),
               "basis": S["reasoning_basis"].format(lvl=it.get("alarm_level")),
               "action": it.get("suggestion", "")} for it in c["high"] + c["medium"]]
    if pr:
        add("table", ([S["col_prio"], S["col_item"], S["col_basis"], S["col_action"]],
                      [[p.get("priority", ""), p.get("item", ""), p.get("basis", ""), p.get("action", "")]
                       for p in sorted(pr, key=lambda p: p.get("priority", "P9"))]))
    elif c["overall"] in (OK, NODATA):
        add("p", S["no_handling"])
    else:
        add("p", S["no_prio"])

    # 5. Remediation
    add("h1", S["h_rem"])
    rem = an.get("remediation") or {}
    for title, key in ((S["h_short"], "short_term"), (S["h_mid"], "medium_term"),
                       (S["h_long"], "long_term")):
        add("h2", title)
        items = rem.get(key) or []
        if items:
            add("bullets", items)
        else:
            add("p", S["no_actions_ok"] if c["overall"] in (OK, NODATA) else S["no_actions"])

    # 6. Conclusion
    add("h1", S["h_concl"])
    con = an.get("conclusion") or {}
    add("p", S["overall_p"].format(rating=R(c["overall"]), text=con.get("overall", "")).strip())
    add("h2", S["h_risks"])
    add("bullets", con.get("key_risks") or [S["none_identified"]])
    add("h2", S["h_trends"])
    add("bullets", con.get("trends") or [S["none"]])
    if an.get("notes"):
        add("h2", S["h_notes"])
        add("bullets", an["notes"])
    return b


# ---------------------------------------------------------------- renderers
def cover_lines(meta, S):
    lines = [(S["cover_site"], meta["site_name"])]
    if meta.get("region"):
        lines.append((S["cover_region"], meta["region"]))
    lines += [(S["cover_time"], meta["inspection_time"]),
              (S["cover_tool"], meta.get("tool") or S["default_tool"])]
    return lines


def md_table(headers, rows):
    esc = lambda s: str(s).replace("|", "\\|").replace("\n", " ")
    out = ["| " + " | ".join(esc(h) for h in headers) + " |", "|" + "---|" * len(headers)]
    out += ["| " + " | ".join(esc(x) for x in r) + " |" for r in rows]
    return out


def render_md(title, cover, blocks):
    out = ["# " + title, ""]
    out += ["- **%s:** %s" % kv for kv in cover] + [""]
    for kind, payload in blocks:
        if kind == "h1":
            out += ["## " + payload, ""]
        elif kind == "h2":
            out += ["### " + payload, ""]
        elif kind == "p":
            out += [payload, ""]
        elif kind == "bullets":
            out += ["- " + x for x in payload] + [""]
        elif kind == "table":
            out += md_table(*payload) + [""]
    return "\n".join(out)


def render_docx(title, cover, blocks, path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    for _ in range(6):
        doc.add_paragraph()
    t = doc.add_heading(title, level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    for label, value in cover:
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.add_run(label + ": ").bold = True
        para.add_run(str(value))
    doc.add_page_break()

    for kind, payload in blocks:
        if kind == "h1":
            doc.add_heading(payload, level=1)
        elif kind == "h2":
            doc.add_heading(payload, level=2)
        elif kind == "p":
            doc.add_paragraph(payload)
        elif kind == "bullets":
            for x in payload:
                doc.add_paragraph(x, style="List Bullet")
        elif kind == "table":
            headers, rows = payload
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = "Table Grid"
            for cell, text in zip(table.rows[0].cells, headers):
                cell.text = ""
                cell.paragraphs[0].add_run(str(text)).bold = True
            for row in rows:
                for cell, text in zip(table.add_row().cells, row):
                    cell.text = str(text)
            doc.add_paragraph()
    doc.save(path)


def render_summary(results, S, stamp):
    """Multi-site summary: one row per site."""
    rows = []
    for data, c, md_path in results:
        rows.append([
            data["meta"]["site_name"],
            S[c["ac_rating"]],
            "%s %s" % (pct(c["access_min"], S["na"]), S[c["access_rating"]]),
            "%s %s" % (pct(c["ap_rate"], S["na"]), S[c["ap_rating"]]),
            "%d / %d" % (len(c["high"]), len(c["medium"])) if c["reasoning_collected"] else S["na"],
            S[c["overall"]],
            os.path.basename(md_path),
        ])
    rows.sort(key=lambda r: r[0])
    out = ["# " + S["summary_title"], "",
           "- **%s:** %s" % (S["sum_generated"], stamp), ""]
    out += md_table([S["sum_col_site"], S["sum_col_ac"], S["sum_col_access"], S["sum_col_ap"],
                     S["sum_col_risk"], S["sum_col_overall"], S["sum_col_report"]], rows)
    return "\n".join(out) + "\n"


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description="Generate the WLAN inspection report (MD + DOCX).")
    ap.add_argument("--data-file", required=True, action="append",
                    help="Structured inspection data JSON (repeat for several sites)")
    ap.add_argument("--output-dir", required=True, help="Directory for the reports")
    ap.add_argument("--lang", choices=sorted(STRINGS), help="Report language (default: meta.language or en)")
    ap.add_argument("--no-docx", action="store_true", help="Write MD only")
    args = ap.parse_args()

    os.makedirs(args.output_dir, exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    docx_missing = False
    results = []
    summary_lang = args.lang

    for path in args.data_file:
        data = load(path)
        lang = args.lang or data["meta"].get("language") or "en"
        if lang not in STRINGS:
            sys.stderr.write("[gen_report] unsupported language '%s'; using en\n" % lang)
            lang = "en"
        summary_lang = summary_lang or lang
        S = STRINGS[lang]
        c = compute(data)
        blocks = build_blocks(data, c, S)
        cover = cover_lines(data["meta"], S)

        site = re.sub(r'[\\/:*?"<>|\s]+', "_", data["meta"]["site_name"]).strip("_")
        base = os.path.join(os.path.abspath(args.output_dir), "Inspection_Report_%s_%s" % (site, stamp))

        md_path = base + ".md"
        with open(md_path, "w", encoding="utf-8") as f:
            f.write(render_md(S["title"], cover, blocks))
        print("MD: " + md_path)
        results.append((data, c, md_path))

        if args.no_docx or docx_missing:
            continue
        try:
            render_docx(S["title"], cover, blocks, base + ".docx")
            print("DOCX: " + base + ".docx")
        except ImportError:
            sys.stderr.write("[gen_report] python-docx is not installed; DOCX skipped. "
                             "Install it with: pip install python-docx\n")
            docx_missing = True

    if len(results) > 1:
        S = STRINGS[summary_lang or "en"]
        sum_path = os.path.join(os.path.abspath(args.output_dir), "Inspection_Summary_%s.md" % stamp)
        with open(sum_path, "w", encoding="utf-8") as f:
            f.write(render_summary(results, S, stamp))
        print("SUMMARY: " + sum_path)

    if docx_missing:
        sys.exit(3)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    main()
